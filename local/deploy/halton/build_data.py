#!/usr/bin/env python3
"""
Готовит данные для import.php из исходников владельца (архив Halton_upload.zip):
  - products.json — товары (название, характеристики, описание, галерея, SEO), тексты
    разделов и страницы бренда;
  - папка с картинками в латинице (JPEG q85; схема KCJ остаётся PNG).

Запуск на хосте (нужны python-docx и Pillow):
  PYTHONPATH=<pylib> python3 build_data.py --src=/tmp/halton/Halton --out=/var/www/prof-equip-test/halton_photos

Всё содержимое карточек берётся из docx без изменений; характеристики — в исходном
порядке, строка делится по ПЕРВОМУ двоеточию, строка без двоеточия идёт без названия.
"""
import argparse
import hashlib
import html
import json
import os
import re
import sys
import unicodedata

import docx
from PIL import Image

HERE = os.path.dirname(os.path.abspath(__file__))

# Схема воздухообмена вентилируемого потолка — лежит и в KCJ (KCJ3.png), и в KVL (KVL2.png).
# К зонту KVL не относится: исключается по хешу, не по имени.
SCHEME_MD5 = '7dcc6ad1cb'

# Порядок галереи (имена после NFC). Правило: файл без номера, затем по возрастанию.
# Исключение — KVI: файл без номера это инфографика, фото зонта — «™2», оно главное.
PRODUCTS = [
    {'folder': 'KCJ', 'slug': 'halton-kcj-capture-jet', 'section': 'ventiliruemye-potolki', 'img': 'halton-kcj-ventiliruemyi-potolok',
     'gallery': ['Вентилируемый потолок HALTON KCJ.png', 'Вентилируемый потолок HALTON KCJ 1.png', 'Вентилируемый потолок HALTON KCJ3.png'],
     'seo_description': 'Вентилируемый потолок HALTON KCJ Capture Jet™ для открытых, производственных и демонстрационных кухонь: вытяжка и подача компенсационного воздуха в одной системе, снижение расхода вытяжного воздуха до 40%, циклонные фильтры HALTON KSA.'},
    {'folder': 'KVF', 'slug': 'halton-kvf-capture-jet', 'section': 'vytyazhnye-zonty', 'img': 'halton-kvf-vytyazhnoi-zont',
     'gallery': ['Вытяжной зонт HALTON KVF Capture Jet™.png', 'Вытяжной зонт HALTON KVF Capture Jet™2.png', 'Вытяжной зонт HALTON KVF Capture Jet™3.png'],
     'seo_description': 'Вытяжной зонт HALTON KVF Capture Jet™ из нержавеющей стали AISI 304 со встроенной фронтальной подачей компенсационного воздуха: снижение расхода вытяжного воздуха до 30–40%, циклонные фильтры HALTON KSA.'},
    {'folder': 'KVI', 'slug': 'halton-kvi-capture-jet', 'section': 'vytyazhnye-zonty', 'img': 'halton-kvi-vytyazhnoi-zont',
     'gallery': ['Вытяжной зонт HALTON KVI Capture Jet™2.png', 'Вытяжной зонт HALTON KVI Capture Jet™.png', 'Вытяжной зонт HALTON KVI Capture Jet™1.png'],
     'seo_description': 'Вытяжной зонт HALTON KVI Capture Jet™ без встроенной подачи компенсационного воздуха: снижение расхода вытяжного воздуха до 30%, высота 600 мм или 450 мм для низких потолков, циклонные фильтры HALTON KSA.'},
    # docx для KVL владелец не прислал — карточка собрана вручную (KVL_MANUAL ниже) из инфографики
    # и схемы в архиве владельца + описания европейской версии KVL у дистрибьютора Halton.
    {'folder': 'KVL', 'slug': 'halton-kvl', 'section': 'vytyazhnye-zonty', 'img': 'halton-kvl-vytyazhnoi-zont', 'manual': True,
     'gallery': ['Вытяжной зонт HALTON KVL.png', 'Вытяжной зонт HALTON KVL1.png', 'Вытяжной зонт HALTON KVL 1.png'],
     'seo_description': 'Пристенный вытяжной зонт HALTON KVL Capture Jet™ с низкой боковой панелью для зон жарки и гриля: снижение расхода вытяжного воздуха до 50%, циклонные фильтры HALTON KSA, LED-освещение до 500 лк.'},
    {'folder': 'KVX', 'slug': 'halton-kvx', 'section': 'vytyazhnye-zonty', 'img': 'halton-kvx-vytyazhnoi-zont',
     'gallery': ['Вытяжной зонт HALTON KVX.png', 'Вытяжной зонт HALTON KVX1.png', 'Вытяжной зонт HALTON KVX2.png'],
     'seo_description': 'Вытяжной зонт HALTON KVX из сатинированной нержавеющей стали AISI 304: циклонные фильтры HALTON KSA, высота 600 мм, наклонное исполнение от 300 до 600 мм для низких потолков. Компенсационный воздух подается отдельно.'},
    {'folder': 'UVF', 'slug': 'halton-uvf-capture-ray', 'section': 'vytyazhnye-zonty', 'img': 'halton-uvf-vytyazhnoi-zont',
     'gallery': ['Вытяжной зонт HALTON UVF Capture Ray™.png', 'Вытяжной зонт HALTON UVF Capture Ray™1.png', 'Вытяжной зонт HALTON UVF Capture Ray™2.png'],
     'seo_description': 'Вытяжной зонт HALTON UVF Capture Ray™ с UV-обработкой вытяжного воздуха и фронтальной подачей компенсационного воздуха: снижение расхода вытяжного воздуха до 30–40%, циклонные фильтры HALTON KSA.'},
    {'folder': 'UVI', 'slug': 'halton-uvi-capture-ray', 'section': 'vytyazhnye-zonty', 'img': 'halton-uvi-vytyazhnoi-zont',
     'gallery': ['Вытяжной зонт HALTON UVI Capture Ray™.png', 'Вытяжной зонт HALTON UVI Capture Ray™1.png', 'Вытяжной зонт HALTON UVI Capture Ray™3.png'],
     'seo_description': 'Вытяжной зонт HALTON UVI Capture Ray™ с UV-обработкой вытяжного воздуха, без встроенной подачи компенсационного воздуха: снижение расхода вытяжного воздуха до 40%, циклонные фильтры HALTON KSA.'},
    {'folder': 'UVL', 'slug': 'halton-uvl-low-sidewall-canopy', 'section': 'vytyazhnye-zonty', 'img': 'halton-uvl-vytyazhnoi-zont',
     'gallery': ['Вытяжной зонт HALTON UVL.png', 'Вытяжной зонт HALTON UVL 2.png'],
     'seo_description': 'Пристенный вытяжной зонт HALTON UVL Low Sidewall Canopy для зон жарки и гриля: технологии Capture Jet™ и Capture Ray™, снижение расхода вытяжного воздуха до 50%, встроенное LED-освещение до 500 лк.'},
]

KVL_MANUAL = {
    'sources': [
        'Архив владельца: KVL/Вытяжной зонт HALTON KVL1.png (инфографика), KVL/Вытяжной зонт HALTON KVL 1.png (схема)',
        'https://stoddart.com.au/products/halton-low-sidewall-extraction-canopy-kvl',
        'https://www.halton.com/app/uploads/2020/08/Halton_CaptureJet_KVL_SS006.pdf',
    ],
    'name': 'Вытяжной зонт HALTON KVL Capture Jet™',
    'specs': [
        ('Тип', 'пристенный вытяжной зонт с низкой боковой панелью'),
        ('Система', 'Capture Jet™'),
        ('Назначение', 'зоны с фритюрницами, грилями и жарочными поверхностями, кухни ресторанов быстрого обслуживания, профессиональные кухни'),
        ('Материал', 'нержавеющая сталь AISI 304'),
        ('Толщина металла', '1,2 мм'),
        ('Конструкция', 'сварная'),
        ('Фильтры', 'циклонные HALTON KSA 500×330×50 мм'),
        ('Эффективность фильтрации', 'до 95% частиц размером 10 мкм и более'),
        ('Сертификация фильтров', 'UL 1046'),
        ('Снижение расхода вытяжного воздуха', 'до 50% по сравнению с обычными пристенными зонтами'),
        ('Система балансировки', 'T.A.B.™'),
        ('Освещение', 'встроенное LED, освещенность до 500 лк'),
        ('Подача компенсационного воздуха', 'опционально — через перфорированную лицевую панель облицовки с низкой скоростью'),
        ('', 'Возможна облицовка с полкой для тарелок'),
        ('', 'Возможна интеграция системы Capture Ray™'),
        ('', 'Возможна установка угольного фильтра'),
        ('', 'Возможна установка встроенной системы пожаротушения'),
        ('Размеры и конфигурация', 'определяются проектом'),
    ],
    'description': [
        'HALTON KVL Capture Jet™ — профессиональный пристенный вытяжной зонт с низкой боковой панелью для кухонь с высокой '
        'тепловой и жировой нагрузкой. Модель предназначена для эффективного захвата и удаления выбросов от фритюрниц, грилей '
        'и жарочных поверхностей на кухнях ресторанов быстрого обслуживания и других профессиональных кухнях.',
        'Сопла Capture Jet™ вдоль нижней части передней кромки зонта формируют вертикальную воздушную завесу. Вместе с вытяжкой '
        'в задней части зонта она удерживает поднимающийся от оборудования загрязненный воздух и не дает ему распространяться '
        'в рабочую зону. Это позволяет снизить требуемый объем вытяжного воздуха <strong>до 50%</strong> по сравнению с обычными '
        'пристенными зонтами. Циклонные фильтры HALTON KSA задерживают до <strong>95% частиц размером 10 мкм и более</strong>.',
        'Корпус сварной конструкции выполнен из нержавеющей стали AISI 304 толщиной 1,2 мм. Встроенные светодиодные светильники '
        'обеспечивают освещенность рабочей зоны до 500 лк, а технология T.A.B.™ упрощает измерение расхода воздуха и '
        'балансировку системы при пусконаладке.',
        'В зависимости от проекта зонт может комплектоваться облицовкой с полкой для тарелок или с перфорированной лицевой '
        'панелью для подачи компенсационного воздуха, угольным фильтром и встроенной системой пожаротушения. Вариант с '
        'UV-обработкой вытяжного воздуха Capture Ray™ — <a href="/product/halton-uvl-low-sidewall-canopy/">HALTON UVL</a>.',
    ],
}

# Фасеты умного фильтра (свойства VENT_*, см. миграцию 2026-09-29-halton-filter-props.php).
# Значения — только то, что прямо сказано в docx / KVL_MANUAL; нет данных — нет значения.
FACETS = {
    'KCJ': {'VENT_KONSTRUKTSIYA': ['Вентилируемый потолок'], 'VENT_TEKHNOLOGII': ['Capture Jet™'],
            'VENT_KOMP_VOZDUKH': ['Встроенная'], 'VENT_SNIZHENIE': ['до 40%'],
            'VENT_OPTSII': ['Встроенная система пожаротушения', 'Система M.A.R.V.E.L.', 'Встроенное освещение']},
    'KVF': {'VENT_KONSTRUKTSIYA': ['Вытяжной зонт'], 'VENT_TEKHNOLOGII': ['Capture Jet™'],
            'VENT_KOMP_VOZDUKH': ['Встроенная'], 'VENT_SNIZHENIE': ['до 30–40%']},
    'KVI': {'VENT_KONSTRUKTSIYA': ['Вытяжной зонт'], 'VENT_TEKHNOLOGII': ['Capture Jet™'],
            'VENT_KOMP_VOZDUKH': ['Отдельной системой вентиляции'], 'VENT_SNIZHENIE': ['до 30%'],
            'VENT_NIZKIE_POTOLKI': ['Есть'], 'VENT_OPTSII': ['Встроенная система пожаротушения']},
    'KVL': {'VENT_KONSTRUKTSIYA': ['Пристенный зонт с низкой боковой панелью'], 'VENT_TEKHNOLOGII': ['Capture Jet™'],
            'VENT_KOMP_VOZDUKH': ['Опционально'], 'VENT_SNIZHENIE': ['до 50%'],
            'VENT_OPTSII': ['Встроенная система пожаротушения', 'Встроенное освещение', 'Угольный фильтр']},
    'KVX': {'VENT_KONSTRUKTSIYA': ['Вытяжной зонт'],
            'VENT_KOMP_VOZDUKH': ['Отдельной системой вентиляции'],
            'VENT_NIZKIE_POTOLKI': ['Есть'], 'VENT_OPTSII': ['Встроенная система пожаротушения']},
    'UVF': {'VENT_KONSTRUKTSIYA': ['Вытяжной зонт'], 'VENT_TEKHNOLOGII': ['Capture Jet™', 'Capture Ray™ (UV-обработка)'],
            'VENT_KOMP_VOZDUKH': ['Встроенная'], 'VENT_SNIZHENIE': ['до 30–40%'], 'VENT_NIZKIE_POTOLKI': ['Есть'],
            'VENT_OPTSII': ['Встроенная система пожаротушения', 'Система M.A.R.V.E.L.', 'Встроенное освещение',
                            'Сенсорная панель Halton Touch Screen', 'Блок очистки воздуха Pollustop']},
    'UVI': {'VENT_KONSTRUKTSIYA': ['Вытяжной зонт'], 'VENT_TEKHNOLOGII': ['Capture Jet™', 'Capture Ray™ (UV-обработка)'],
            'VENT_KOMP_VOZDUKH': ['Отдельной системой вентиляции'], 'VENT_SNIZHENIE': ['до 40%'], 'VENT_NIZKIE_POTOLKI': ['Есть'],
            'VENT_OPTSII': ['Встроенная система пожаротушения', 'Система M.A.R.V.E.L.',
                            'Сенсорная панель Halton Touch Screen', 'Блок очистки воздуха Pollustop']},
    'UVL': {'VENT_KONSTRUKTSIYA': ['Пристенный зонт с низкой боковой панелью'], 'VENT_TEKHNOLOGII': ['Capture Jet™', 'Capture Ray™ (UV-обработка)'],
            'VENT_KOMP_VOZDUKH': ['Опционально'], 'VENT_SNIZHENIE': ['до 50%'],
            'VENT_OPTSII': ['Встроенная система пожаротушения', 'Встроенное освещение', 'Угольный фильтр']},
}

SECTIONS = {
    'resheniya-dlya-ventilyatsii': {
        'meta_title': 'Решения для вентиляции профессиональной кухни — купить в ПРОФЭКВИП',
        'meta_description': 'Вытяжные зонты и вентилируемые потолки для ресторанов, гостиниц и профессиональных кухонь: технологии Capture Jet™ и Capture Ray™, циклонные фильтры. Подбор и заказ в ПРОФЭКВИП.',
        'picture': ('UVF', 'Вытяжной зонт HALTON UVF Capture Ray™.png'),
        'description_html': (
            '<h2>Решения для вентиляции профессиональной кухни</h2>\n'
            '<p>Вентиляция профессиональной кухни удаляет тепло, пар, жировые частицы и запахи, обеспечивая комфортные '
            'условия работы персонала и безопасность объекта. В разделе собраны решения для кухонь ресторанов, гостиниц '
            'и других объектов HoReCa:</p>\n'
            '<ul>\n'
            '    <li><a href="/product-category/vytyazhnye-zonty/">Вытяжные зонты</a> — устанавливаются над тепловым '
            'оборудованием и удаляют тепло, пар, жировые аэрозоли и продукты приготовления пищи;</li>\n'
            '    <li><a href="/product-category/ventiliruemye-potolki/">Вентилируемые потолки</a> — объединяют вытяжку и '
            'подачу компенсационного воздуха в одной потолочной системе для открытых, производственных и '
            'демонстрационных кухонь.</li>\n'
            '</ul>\n'
            '<p>Модели можно отфильтровать по конструкции, технологии, способу подачи компенсационного воздуха и '
            'опциям. ПРОФЭКВИП поставляет оборудование для вентиляции в рамках комплексного оснащения гостиниц, '
            'ресторанов и профессиональных кухонь — оставьте заявку на сайте, и специалисты помогут с подбором.</p>'
        ),
    },
    'ventiliruemye-potolki': {
        'meta_title': 'Вентилируемые потолки для профессиональной кухни — купить в ПРОФЭКВИП',
        'meta_description': 'Вентилируемые потолки для открытых, производственных и демонстрационных кухонь: вытяжка и подача компенсационного воздуха в одной системе. Подбор и заказ в ПРОФЭКВИП.',
        'picture': ('KCJ', 'Вентилируемый потолок HALTON KCJ.png'),
        'description_html': (
            '<h2>Вентилируемые потолки для профессиональной кухни</h2>\n'
            '<p>Вентилируемый потолок объединяет вытяжку и подачу компенсационного воздуха в одной потолочной системе: '
            'тепло, пар и загрязненный воздух захватываются непосредственно над рабочей зоной кухни, а свежий воздух '
            'подается с низкой скоростью через лицевую часть потолка. Такие системы применяют на открытых, '
            'производственных и демонстрационных кухнях — там, где важны открытая архитектура кухни, равномерное '
            'распределение воздуха, высокий уровень гигиены и аккуратный внешний вид инженерных систем.</p>\n'
            '<p>Если над оборудованием нужен отдельный зонт, смотрите раздел '
            '<a href="/product-category/vytyazhnye-zonty/">вытяжные зонты</a>. Специалисты ПРОФЭКВИП помогут подобрать '
            'решение под проект — оставьте заявку на сайте.</p>'
        ),
    },
    'vytyazhnye-zonty': {
        'meta_title': 'Вытяжные зонты для профессиональной кухни — купить в ПРОФЭКВИП',
        'meta_description': 'Профессиональные вытяжные зонты для ресторанов, гостиниц, открытых и демонстрационных кухонь: технологии Capture Jet™ и Capture Ray™, циклонные фильтры. Подбор и заказ в ПРОФЭКВИП.',
        'picture': ('KVF', 'Вытяжной зонт HALTON KVF Capture Jet™.png'),
        'description_html': (
            '<h2>Вытяжные зонты для профессиональной кухни</h2>\n'
            '<p>Вытяжной зонт удаляет тепло, пар, жировые аэрозоли и продукты приготовления пищи над тепловым '
            'оборудованием. В каталоге ПРОФЭКВИП представлены профессиональные вытяжные зонты для ресторанов, гостиниц, '
            'открытых и демонстрационных кухонь.</p>\n'
            '<h3>На что обратить внимание при выборе</h3>\n'
            '<ul>\n'
            '    <li><strong>Подача компенсационного воздуха.</strong> Встроенная — через фронтальную часть зонта — или отдельной системой вентиляции помещения;</li>\n'
            '    <li><strong>Технология Capture Jet™.</strong> Направленные воздушные струи удерживают тепловой поток в зоне вытяжки и позволяют снизить необходимый объем вытяжного воздуха;</li>\n'
            '    <li><strong>UV-обработка Capture Ray™.</strong> Нейтрализует жировые пары и снижает запахи в вытяжном воздухе — особенно важно для длинных и горизонтальных воздуховодов;</li>\n'
            '    <li><strong>Фильтрация.</strong> Циклонные фильтры задерживают жировые частицы и уменьшают их накопление в воздуховодах и вытяжном вентиляторе;</li>\n'
            '    <li><strong>Высота помещения.</strong> Для помещений с низкими потолками предусмотрены исполнения меньшей высоты.</li>\n'
            '</ul>\n'
            '<p>Для открытых кухонь смотрите также <a href="/product-category/ventiliruemye-potolki/">вентилируемые потолки</a>. '
            'Не уверены, какое решение нужно? Оставьте заявку на сайте — специалисты ПРОФЭКВИП помогут с подбором.</p>'
        ),
    },
}

BRAND_SEO_DESCRIPTION = ('Halton — международный бренд профессиональных систем вентиляции для гостиниц, ресторанов и '
                         'коммерческих кухонь: вытяжные зонты, вентиляционные потолки, системы очистки воздуха. '
                         'Поставки оборудования Halton — в ПРОФЭКВИП.')


def nfc(s):
    return unicodedata.normalize('NFC', s)


def md5(path):
    return hashlib.md5(open(path, 'rb').read()).hexdigest()


def runs_html(paragraph):
    """Текст абзаца в HTML: соседние runs сшиваются, жирные фрагменты -> <strong>,
    пробелы по краям жирного выносятся наружу (артефакты вида «Capture**** ****Jet»)."""
    groups = []
    for r in paragraph.runs:
        bold = bool(r.bold)
        if groups and groups[-1][0] == bold:
            groups[-1][1] += r.text
        else:
            groups.append([bold, r.text])
    out = ''
    for bold, text in groups:
        if bold and text.strip():
            lead = text[:len(text) - len(text.lstrip())]
            trail = text[len(text.rstrip()):]
            out += html.escape(lead) + '<strong>' + html.escape(text.strip()) + '</strong>' + html.escape(trail)
        else:
            out += html.escape(text)
    return re.sub(r'\s+', ' ', out).strip()


def plain(paragraph):
    return re.sub(r'\s+', ' ', paragraph.text).strip()


def parse_product_docx(path):
    paras = docx.Document(path).paragraphs
    name, specs, desc = None, [], []
    mode = None
    for p in paras:
        text = p.text.strip()
        if not text:
            continue
        if text.startswith('Название:'):
            name = re.sub(r'\s+', ' ', text[len('Название:'):]).strip()
            mode = 'name'
            continue
        if text.startswith('Технические характеристики:'):
            mode = 'specs'
            continue
        if text.startswith('Описание:'):
            mode = 'desc'
            # первый абзац описания в том же параграфе после перевода строки
            h = runs_html(p)
            h = re.sub(r'^(<strong>)?Описание:(</strong>)?\s*', '', h)
            if h:
                desc.append(h)
            continue
        if mode == 'specs':
            line = plain(p)
            if ':' in line:
                k, v = line.split(':', 1)
                specs.append({'name': k.strip(), 'value': v.strip()})
            else:
                specs.append({'name': '', 'value': line})
        elif mode == 'desc':
            desc.append(runs_html(p))
    if not name or not specs or not desc:
        raise SystemExit(f'{path}: не распознана структура (name={name!r}, specs={len(specs)}, desc={len(desc)})')
    return name, specs, desc


def parse_brand_docx(path):
    paras = [p for p in docx.Document(path).paragraphs if p.text.strip()]
    blocks = []
    for i, p in enumerate(paras):
        all_bold = all(r.bold for r in p.runs if r.text.strip())
        text = re.sub(r'\s+', ' ', p.text).strip()
        if all_bold:
            blocks.append({'type': 'h1' if i == 0 else 'h2', 'text': text})
        else:
            blocks.append({'type': 'p', 'html': runs_html(p)})
    return blocks


def convert(src, dst_base, out_dir):
    """PNG -> JPEG q85 (<= 400 КБ, без обрезки/изменения пропорций). Схема KCJ (маленький PNG
    с линиями) остаётся PNG — там JPEG только испортит качество."""
    im = Image.open(src)
    if md5(src).startswith(SCHEME_MD5):
        dst = os.path.join(out_dir, dst_base + '.png')
        im.save(dst, optimize=True)
        return dst
    dst = os.path.join(out_dir, dst_base + '.jpg')
    im = im.convert('RGB')
    for q in (85, 80, 75):
        im.save(dst, 'JPEG', quality=q, optimize=True, progressive=True)
        if os.path.getsize(dst) <= 400 * 1024:
            break
    return dst


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--src', required=True)
    ap.add_argument('--out', required=True)
    a = ap.parse_args()
    os.makedirs(a.out, exist_ok=True)

    listing = {}
    for d, _, files in os.walk(a.src):
        for f in files:
            if f == '.DS_Store' or f.startswith('._') or '__MACOSX' in d:
                continue
            listing[(nfc(os.path.relpath(d, a.src)), nfc(f))] = os.path.join(d, f)

    products = []
    for idx, spec in enumerate(PRODUCTS, 1):
        folder = spec['folder']
        docs = [p for (d, f), p in listing.items() if d == folder and f.endswith('.docx')]
        if spec.get('manual'):
            if docs:
                raise SystemExit(f'{folder}: появился docx владельца — уберите manual и перегенерируйте из него')
            name = KVL_MANUAL['name']
            specs = [{'name': k, 'value': v} for k, v in KVL_MANUAL['specs']]
            desc = KVL_MANUAL['description']
        else:
            if len(docs) != 1:
                raise SystemExit(f'{folder}: ожидался 1 docx, найдено {len(docs)}')
            name, specs, desc = parse_product_docx(docs[0])

        folder_pngs = {f for (d, f) in listing if d == folder and f.endswith('.png')}
        unused = folder_pngs - set(spec['gallery'])
        images = []
        for n, fname in enumerate(spec['gallery'], 1):
            src = listing.get((folder, fname))
            if not src:
                raise SystemExit(f'{folder}: нет файла {fname}')
            is_scheme = md5(src).startswith(SCHEME_MD5)
            if is_scheme and folder != 'KCJ':
                raise SystemExit(f'{folder}: схема KCJ попала в галерею')
            dst = convert(src, f"{spec['img']}-{n}", a.out)
            alt = ('Схема воздухообмена вентилируемого потолка HALTON KCJ Capture Jet™' if is_scheme
                   else f'{name} — фото {n}')
            images.append({'file': os.path.basename(dst), 'source': fname, 'alt': alt,
                           'bytes': os.path.getsize(dst)})
        for fname in sorted(unused):
            is_scheme = md5(listing[(folder, fname)]).startswith(SCHEME_MD5)
            print(f'{folder}: исключён {fname}' + (' (дубликат схемы KCJ)' if is_scheme else ''))
            if not is_scheme:
                raise SystemExit(f'{folder}: неучтённый файл {fname}')

        products.append({
            'folder': folder,
            'slug': spec['slug'],
            'section': spec['section'],
            'active': not spec.get('draft', False),
            'sort': 500 + idx,
            'name': name,
            'sources': KVL_MANUAL['sources'] if spec.get('manual') else ['docx владельца'],
            'specs': specs,
            'facets': FACETS[folder],
            'description_paragraphs': desc,
            'description_html': '\n'.join(f'<p>{p}</p>' for p in desc),
            'preview_text': re.sub(r'<[^>]+>', '', desc[0]) if desc else '',
            'images': images,
            'seo': {'title': f'{name} — купить в ПРОФЭКВИП', 'description': spec['seo_description']},
        })
        print(f"{folder}: «{name}», характеристик {len(specs)}, абзацев {len(desc)}, фото {len(images)}")

    brand_blocks = parse_brand_docx(listing[('.', 'Halton.docx')])
    h1 = next(b['text'] for b in brand_blocks if b['type'] == 'h1')
    # H1 страницы бренда шаблон всегда строит из названия элемента («Halton»), поэтому заголовок
    # из docx идёт первым <h2> текста и в SEO title.
    brand_html = '\n'.join(
        f"<h2>{html.escape(b['text'])}</h2>" if b['type'] in ('h1', 'h2') else f"<p>{b['html']}</p>"
        for b in brand_blocks)

    sections = {}
    for code, s in SECTIONS.items():
        s = dict(s)
        if s['picture']:
            folder, fname = s['picture']
            s['picture'] = os.path.basename(convert(listing[(folder, fname)], f'{code}', a.out))
        sections[code] = s

    data = {
        'source': 'Halton_upload.zip от владельца, 2026-09-29',
        'brand': {'name': 'Halton', 'code': 'halton', 'h1': h1, 'text_html': brand_html,
                  'meta_title': h1, 'meta_description': BRAND_SEO_DESCRIPTION},
        'sections': sections,
        'products': products,
    }
    with open(os.path.join(HERE, 'products.json'), 'w', encoding='utf-8') as fh:
        json.dump(data, fh, ensure_ascii=False, indent=2)
    print('products.json записан')


if __name__ == '__main__':
    sys.exit(main())
